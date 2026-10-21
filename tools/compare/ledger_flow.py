"""Runs the Payments scenarios (add with bills, validation errors, edit, bill
download, delete, Excel) against the Django and the Go site and compares the
answers and the rows left behind.

    ledger_flow.py
"""
import io
import os
import re
import sys
import uuid

sys.path.insert(0, os.path.dirname(__file__))
from pages import Site, norm  # noqa: E402

import psycopg2  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

DB = dict(host="127.0.0.1", port=5433, dbname="deedflow", user="deedflow", password="deedflow")


def multipart(fields, files):
    boundary = "----df" + uuid.uuid4().hex
    out = io.BytesIO()
    for name, value in fields:
        out.write(f'--{boundary}\r\nContent-Disposition: form-data; name="{name}"\r\n\r\n{value}\r\n'.encode())
    for name, filename, content, ctype in files:
        out.write(f'--{boundary}\r\nContent-Disposition: form-data; name="{name}"; filename="{filename}"\r\n'
                  f'Content-Type: {ctype}\r\n\r\n'.encode())
        out.write(content + b"\r\n")
    out.write(f"--{boundary}--\r\n".encode())
    return out.getvalue(), f"multipart/form-data; boundary={boundary}"


def post(site, path, fields, files=()):
    fields = list(fields) + [("csrfmiddlewaretoken", site.csrf())]
    body, ctype = multipart(fields, files)
    return site.req("POST", path, raw_body=body, content_type=ctype)


def alerts(html):
    return [re.sub(r"\s+", " ", m).strip() for m in re.findall(r'<div>(.*?)</div>\s*<button type="button" class="btn-close"', html, re.S)]


def errors(html):
    return sorted(set(re.sub(r"\s+", " ", m).strip() for m in re.findall(r'class="text-danger small mt-1">(.*?)</div>', html, re.S)))


def rows(prefix):
    with psycopg2.connect(**DB) as c, c.cursor() as cur:
        cur.execute("select e.id, e.date, e.category, e.amount, e.payment_mode, e.description, e.created_by_id, "
                    "(select array_agg(b.original_name order by b.id) from ledger_expensebill b where b.expense_id=e.id), "
                    "(select array_agg(b.file order by b.id) from ledger_expensebill b where b.expense_id=e.id) "
                    "from ledger_expense e where e.organization_id=1 and e.description like %s order by e.id", (prefix + "%",))
        return cur.fetchall()


def norm_rows(rs, prefix):
    out = []
    for r in rs:
        files = r[8] or []
        shape = [re.sub(r"[0-9a-f]{32}", "<hex>", f) for f in files]
        out.append((str(r[1]), r[2], str(r[3]), r[4], r[5].replace(prefix, "P"), r[6], r[7], shape))
    return out


def run(site, tag):
    site.login()
    log = []
    desc = lambda s: f"{tag} {s}"  # noqa: E731

    def step(name, resp, follow=False):
        s, h, b = resp
        entry = [name, s, h.get("Location")]
        if s == 200:
            entry += [errors(b)]
        if s == 302 and follow:
            s2, _, b2 = site.req("GET", h.get("Location"))
            entry += [s2, alerts(b2)]
        log.append(entry)
        return s, h, b

    photo = ("bills", "IMG_001.jpg", b"\xff\xd8\xff fake jpeg", "image/jpeg")
    pdf = ("bills", "bill.pdf", b"%PDF-1.4 fake", "application/pdf")
    base = [("date", "2026-10-01"), ("category", "rent"), ("amount", "8000"), ("payment_mode", "upi"),
            ("description", desc("October rent"))]
    step("add with bills", post(site, "/accounts/payments/add/", base, [photo, pdf]), follow=True)
    step("add and another", post(site, "/accounts/payments/add/", base[:4] + [("description", desc("second")), ("next", "add")]))
    step("bad extension", post(site, "/accounts/payments/add/", base[:4] + [("description", desc("bad"))],
                               [("bills", "notes.exe", b"MZ", "application/octet-stream")]))
    step("zero amount", post(site, "/accounts/payments/add/", [("date", "2026-10-01"), ("category", "rent"), ("amount", "0"),
                                                             ("description", desc("zero"))]))
    step("missing fields", post(site, "/accounts/payments/add/", [("description", desc("missing"))]))
    step("bad amount", post(site, "/accounts/payments/add/", [("date", "2026-10-01"), ("category", "rent"), ("amount", "12345678901.234"),
                                                             ("description", desc("digits"))]))
    step("bad category", post(site, "/accounts/payments/add/", [("date", "2026-10-01"), ("category", "nope"), ("amount", "5"),
                                                              ("description", desc("cat"))]))
    step("bad date", post(site, "/accounts/payments/add/", [("date", "31/02/2026"), ("category", "rent"), ("amount", "5"),
                                                          ("description", desc("date"))]))
    step("too many bills", post(site, "/accounts/payments/add/", base[:4] + [("description", desc("many"))],
                                [("bills", f"a{i}.png", b"png", "image/png") for i in range(11)]))
    mine = [r for r in rows(tag) if r[5] == desc("October rent")]
    pk = mine[0][0]
    bills = psycopg2.connect(**DB)
    with bills, bills.cursor() as cur:
        cur.execute("select id, original_name from ledger_expensebill where expense_id=%s order by id", (pk,))
        bill_rows = cur.fetchall()
    jpg_id = [b[0] for b in bill_rows if b[1] == "IMG_001.jpg"][0]
    pdf_id = [b[0] for b in bill_rows if b[1] == "bill.pdf"][0]
    for name, query in (("inline", ""), ("download", "?download=1")):
        s, h, b = site.req("GET", f"/accounts/payments/bills/{pdf_id}/{query}")
        log.append([f"bill {name}", s, h.get("Content-Type"), h.get("Content-Disposition"), h.get("Cache-Control"), b[:12]])
    s, h, b = site.req("GET", f"/accounts/payments/bills/{jpg_id}/")
    log.append(["bill jpg", s, h.get("Content-Type"), h.get("Content-Disposition")])
    step("edit", post(site, f"/accounts/payments/{pk}/", [("date", "2026-10-01"), ("category", "rent"), ("amount", "8500"),
                                                       ("description", desc("October rent")), ("remove_bill", str(jpg_id))],
                      [("bills", "new.png", b"png", "image/png")]), follow=True)
    step("delete via GET", site.req("GET", f"/accounts/payments/{pk}/delete/"))
    s, h, b = site.req("GET", "/accounts/payments/export/?q=" + tag.replace(" ", "+"), binary=True)
    wb = load_workbook(io.BytesIO(b))
    log.append(["export", s, h.get("Content-Type"), wb.sheetnames, [[(c.value if not isinstance(c.value, str) else c.value.replace(tag, "P")) for c in row] for row in wb["Payments"].iter_rows()],
                [wb["Payments"].cell(row=2, column=i).number_format for i in (1, 5)]])
    s, h, b = site.req("GET", "/accounts/export/", binary=True)
    wb = load_workbook(io.BytesIO(b))
    log.append(["accounts export", s, wb.sheetnames, [[c.value for c in row] for row in wb["Stamp payment"].iter_rows()],
                [[c.value for c in row] for row in wb["Fee payment"].iter_rows()],
                [wb["Stamp payment"].cell(row=2, column=i).number_format for i in (4, 8, 12, 15)]])
    step("delete", post(site, f"/accounts/payments/{pk}/delete/", []), follow=True)
    for r in rows(tag):
        pass
    step("collect (no portal login)", post(site, "/accounts/collect/", []), follow=True)
    s, h, b = site.req("GET", "/accounts/collect/")
    log.append(["collect GET", s])
    return log


def main():
    dj = Site("http://127.0.0.1:8091", "admin", "correct-horse-7")
    go = Site("http://127.0.0.1:8090", "admin", "correct-horse-7")
    with psycopg2.connect(**DB) as c, c.cursor() as cur:
        cur.execute("delete from ledger_expensebill where expense_id in (select id from ledger_expense where organization_id=1 "
                    "and (description like 'DJ%' or description like 'GO%'))")
        cur.execute("delete from ledger_expense where organization_id=1 and (description like 'DJ%' or description like 'GO%')")
        cur.execute("update accounts_portalcredential set portal_username='' where organization_id=1")
    a, b = run(dj, "DJ"), run(go, "GO")
    ok = True
    for x, y in zip(a, b):
        if x != y:
            ok = False
            print("DIFFERENT:", x[0])
            print("  django:", x[1:])
            print("  go:    ", y[1:])
        else:
            print("same:", x[0])
    ra, rb = norm_rows(rows("DJ"), "DJ"), norm_rows(rows("GO"), "GO")
    print("rows", "same" if ra == rb else f"DIFFERENT\n  {ra}\n  {rb}")
    sys.exit(0 if ok and ra == rb else 1)


main()
